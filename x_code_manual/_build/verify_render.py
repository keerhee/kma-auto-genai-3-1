# -*- coding: utf-8 -*-
"""최종 검수 — 실제로 그려진 글자와 실제 구분선 위치를 맞대어 본다 (§1.6).

추정 기반 qa_overflow.py가 못 잡는 결함을 잡는다. 구분선은 그림이라 텍스트
검사에 안 걸리고, 글상자끼리는 서로 안 겹친 채 선만 관통하는 일이 흔하다.

사용자가 본 증상은 두 가지다.
 ① 가로 구분선이 글자를 가로지른다  ② 아랫줄 글자와 윗줄 글자가 겹친다
①은 PDF 글자 상자만 봐서는 안 보인다(선은 그림이라 텍스트로 안 잡힌다).
그래서 선 좌표는 PPTX에서 정확히 가져오고, 글자 상자는 렌더 결과에서 잰다.

사용: python3 scripts/verify_render.py <deck.pptx> [rendered.pdf]
      pdf를 안 주면 LibreOffice로 그 자리에서 렌더한다.
의존: python-pptx, poppler(pdftotext), LibreOffice
"""
import re
import subprocess
import sys
from collections import defaultdict

from pptx import Presentation

EMU_PT = 12700.0
GAP_SPLIT = 14.0     # 가로로 이만큼 벌어지면 다른 조각(나란한 카드 구분)
OVERLAP_FRAC = 0.34  # 줄 높이의 이 비율을 넘게 물릴 때만 겹침
ROW_TOL = 1.5        # 이 정도 어긋난 yMin은 같은 줄로 본다
RULE_PAD = 1.0       # 선이 글자 상자 안쪽으로 이만큼 넘어오면 가로지른 것


def is_pictograph(t):
    """문장부호·기호·이모지만으로 된 조각.

    이런 글자는 글자 상자가 본문 줄보다 높거나 낮게 잡혀(폰트 대체 때문에)
    같은 줄의 옆 글자와 늘 스친다. 눈으로는 한 줄인데 검사에서만 겹친다.
    예: "[지시] ①~⑤"의 물결표, "웹 · 앱"의 가운뎃점.
    """
    import unicodedata
    core = [c for c in t if not c.isspace()]
    return bool(core) and all(
        unicodedata.category(c)[0] in ("P", "S") or
        (ord(c) > 0x2000 and not c.isalnum() and not ("\uac00" <= c <= "\ud7a3"))
        for c in core)


def layered(a, b):
    """일부러 겹쳐 놓은 것 — 큰 장식 글자 위에 작은 라벨을 얹은 경우."""
    ha, hb = a[1] - a[0], b[1] - b[0]
    if max(ha, hb) < 2 * min(ha, hb):
        return False
    big, small = (a, b) if ha >= hb else (b, a)
    return (big[0] <= small[0] + 2 and big[1] >= small[1] - 2
            and big[2] <= small[2] + 2 and big[3] >= small[3] - 2)


def fragments(pg):
    # 같은 줄인데 글꼴이 다르면(숫자·기호가 대체 글꼴로 그려질 때) yMin이 0.1pt
    # 안팎으로 어긋난다. 정확히 같은 값으로 묶으면 한 줄이 두 조각으로 갈라져
    # 서로 "겹쳤다"고 잡힌다. 그래서 허용오차를 두고 묶는다.
    words = [(float(ymin), float(xmin), float(xmax), float(ymax), w)
             for xmin, ymin, xmax, ymax, w in re.findall(
                 r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" '
                 r'yMax="([\d.]+)">(.*?)</word>', pg)]
    rows, cur, base = [], [], None
    for ymin, xmin, xmax, ymax, w in sorted(words):
        if base is None or ymin - base <= ROW_TOL:
            if base is None:
                base = ymin
            cur.append((xmin, xmax, ymin, ymax, w))
        else:
            rows.append(cur)
            cur, base = [(xmin, xmax, ymin, ymax, w)], ymin
    if cur:
        rows.append(cur)
    out = []
    for ws in rows:
        ws = sorted(ws)
        cur = [ws[0]]
        for w in ws[1:]:
            if w[0] - cur[-1][1] > GAP_SPLIT:
                out.append(cur)
                cur = [w]
            else:
                cur.append(w)
        out.append(cur)
    return [(min(w[2] for w in f), max(w[3] for w in f),
             min(w[0] for w in f), max(w[1] for w in f),
             "".join(w[4] for w in f)) for f in out]


def _opaque(sh):
    """뒤 그림을 가리는 불투명 도형인가 (단색 채움)."""
    try:
        return sh.fill.type is not None and sh.fill.fore_color.rgb is not None
    except Exception:
        return False


def rules(prs):
    """{쪽번호: [(y, x0, x1, 덮개목록)]} — PPTX의 0높이 구분선 좌표(pt).

    선 뒤에(=z 순서상 나중에) 그려진 불투명 도형은 그 선을 가린다. 가려진
    선은 화면에 없으므로 글자를 가로질러도 결함이 아니다. 덮개 사각형을
    함께 들고 다니며 판정할 때 걸러낸다.
    """
    out = {}
    for i, s in enumerate(prs.slides, 1):
        shapes = list(s.shapes)
        rs = []
        for z, sh in enumerate(shapes):
            if sh.height != 0 or sh.left is None or not sh.width:
                continue
            covers = [(o.left / EMU_PT, (o.left + o.width) / EMU_PT,
                       o.top / EMU_PT, (o.top + o.height) / EMU_PT)
                      for o in shapes[z + 1:]
                      if o.left is not None and o.height and _opaque(o)]
            rs.append((sh.top / EMU_PT, sh.left / EMU_PT,
                       (sh.left + sh.width) / EMU_PT, covers))
        out[i] = rs
    return out


def hidden(ry, x0, x1, covers):
    """글자가 놓인 x 구간에서 이 선이 덮개에 가려지는가."""
    return any(cx0 <= x0 + 1 and cx1 >= x1 - 1 and cy0 <= ry <= cy1
               for cx0, cx1, cy0, cy1 in covers)


def main(pptx, pdf, show=30):
    prs = Presentation(pptx)
    rl = rules(prs)
    xml = subprocess.run(["pdftotext", "-bbox", pdf, "-"],
                         capture_output=True, text=True).stdout
    hits = []
    for pno, pg in enumerate(re.split(r"<page ", xml)[1:], 1):
        fr = sorted(fragments(pg))
        for y0, y1, x0, x1, t in fr:
            for ry, rx0, rx1, covers in rl.get(pno, []):
                if hidden(ry, x0, x1, covers):
                    continue
                if y0 + RULE_PAD < ry < y1 - RULE_PAD and x0 < rx1 and x1 > rx0:
                    hits.append((pno, "선이 글자를 가로지름",
                                 min(ry - y0, y1 - ry), t[:40], ""))
                    break
        for a in range(len(fr)):
            for b in range(a + 1, len(fr)):
                ay0, ay1, ax0, ax1, at = fr[a]
                by0, by1, bx0, bx1, bt = fr[b]
                if by0 >= ay1:
                    break
                ov = min(ay1, by1) - max(ay0, by0)
                if ov <= OVERLAP_FRAC * min(ay1 - ay0, by1 - by0):
                    continue
                if ax0 >= bx1 - 1 or bx0 >= ax1 - 1:
                    continue
                if is_pictograph(at) or is_pictograph(bt):
                    continue
                if layered((ay0, ay1, ax0, ax1), (by0, by1, bx0, bx1)):
                    continue
                hits.append((pno, "글자끼리 겹침", ov, at[:36], bt[:36]))
    for pno, kind, amt, a, b in hits[:show]:
        tail = f"  ↕  {b}" if b else ""
        print(f"  p{pno:>3} {kind} {amt:5.1f}pt  {a}{tail}")
    if len(hits) > show:
        print(f"  … 외 {len(hits) - show}건")
    print(f"  렌더 겹침 {len(hits)}건")
    return len(hits)


def render(pptx):
    import subprocess as sp
    import tempfile
    out = tempfile.mkdtemp()
    sp.run(["/Applications/LibreOffice.app/Contents/MacOS/soffice", "--headless",
            "--convert-to", "pdf", "--outdir", out, pptx],
           capture_output=True, check=True)
    return os.path.join(out, os.path.splitext(os.path.basename(pptx))[0] + ".pdf")


if __name__ == "__main__":
    import os
    deck = sys.argv[1]
    pdf = sys.argv[2] if len(sys.argv) > 2 else render(deck)
    sys.exit(1 if main(deck, pdf) else 0)
