"""Day2 공통 도구 — xlsx(엑셀 내장 차트)와 단일 HTML(인라인 SVG) 생성.

규칙은 CLAUDE.md에 있다. 이 모듈은 그 규칙을 코드로 강제한다.
- 막대·누적 막대의 y축은 0에서 시작한다.
- 차트 아래에 기준·기간·출처 노트를 붙인다.
- HTML은 외부 리소스를 쓰지 않는다.
"""
from __future__ import annotations

import base64
import html
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# 윈도우 콘솔에서도 한글 메시지(특히 [중단] 메시지)가 깨지지 않게 한다
for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8")
    except Exception:
        pass

ROOT = Path(__file__).resolve().parents[2]
INPUT = ROOT / "input"
WORKING = ROOT / "working"
OUTPUT = ROOT / "output"

ACCENT = "1F5FAD"   # 강조 1색
MUTED = "A7B0BC"    # 나머지 회색
NEG = "C0504D"      # 부정(감정분석 전용)
NEU = "BFBFBF"      # 중립


def read_csv(name: str) -> pd.DataFrame:
    """input/ 또는 working/의 CSV를 UTF-8(BOM) 기준으로 읽는다."""
    p = Path(name)
    if not p.is_absolute():
        # 'working/x.csv'처럼 폴더를 붙여 쓰면 프로젝트 기준, 파일 이름만 쓰면 input → working 순
        p = next((c for c in (ROOT / name, INPUT / name, WORKING / name) if c.exists()), WORKING / name)
    return pd.read_csv(p, encoding="utf-8-sig")


# ───────────────────────── 입력 선택 · 판단 기록 ─────────────────────────
import hashlib  # noqa: E402
import json  # noqa: E402
from datetime import datetime  # noqa: E402

TOP5_INPUT = INPUT / "Day2_국내EV판매_Top5_실데이터.csv"
STAMPS = WORKING / "_판단기록.json"
CLINIC_CFG = WORKING / "차트클리닉_설정.json"


def top5_source() -> Path:
    """Top5 입력 = input 백업과 working 수집본 중 **가장 최근에 바뀐 파일**.

    수집본이 있다고 무조건 쓰면 input을 새로 넣어도 반영되지 않는다(수정 1).
    """
    cands = [TOP5_INPUT] + sorted(WORKING.glob("수집_EV_Top5_*.csv"))
    return max((p for p in cands if p.exists()), key=lambda p: p.stat().st_mtime)


def top5_table() -> tuple[pd.DataFrame, str, str, Path]:
    """Top5 표를 읽어 (표, 수량 열 이름, 기준 이름, 원천) 을 돌려준다.

    수량 열 이름과 기준은 집계기준 열에서 정한다 — 문구를 코드에 고정하지 않는다(수정 3).
    기준이 행마다 섞여 있으면 멈춘다(CLAUDE.md §2).
    """
    src = top5_source()
    df = read_csv(str(src))
    basis_col = "집계기준(판매/등록)"
    kinds = set()
    for b in df[basis_col].astype(str):
        # 괄호 안 설명('기사 표기 판매순위' 등)은 빼고 앞부분으로 판정한다
        head = b.split("(")[0]
        kinds.add("등록" if "등록" in head else "판매" if ("판매" in head or "출고" in head) else "불명")
    if len(kinds) != 1 or "불명" in kinds:
        raise SystemExit(f"[중단] {src.name}: 집계기준이 하나로 정해지지 않음 {sorted(kinds)} — 사람 확인 필요")
    basis = kinds.pop()
    qty = "등록대수(대)" if basis == "등록" else "판매량(대)"
    df = df.rename(columns={"판매량(대)": qty, basis_col: "집계기준"})
    df = df[["순위", "모델명", "제조사", qty, "기준기간", "집계기준", "출처URL"]]
    return df, qty, basis, src


def market_tables() -> dict:
    """연간 제조사별·연도별 추이 표와 제목·노트를 데이터에서 만든다(문구 고정 금지, 수정 3)."""
    mk = read_csv("Day2_국내EV_월별또는브랜드점유율_실데이터.csv")
    g = mk[mk["구분"].str.contains("제조사별")]
    label = g["구분"].iloc[0].split()[0]                      # 예: '2025'
    basis = "/".join(dict.fromkeys(g["집계기준(판매/등록)"].astype(str)))
    org = "/".join(dict.fromkeys(g["원자료기관"].astype(str)))
    total_row = g[g["항목"].str.startswith("전체")]
    total = int(total_row["값"].iloc[0])
    makers = g[~g["항목"].str.startswith("전체")][["항목", "값", "기준기간", "집계기준(판매/등록)", "원자료기관", "출처URL"]]
    makers = makers.rename(columns={"항목": "제조사", "값": f"{basis}(대)", "집계기준(판매/등록)": "집계기준"}).reset_index(drop=True)

    y = mk[mk["구분"].str.startswith("연도별")]
    ybasis = "/".join(dict.fromkeys(y["집계기준(판매/등록)"].astype(str)))
    yorg = "/".join(dict.fromkeys(y["원자료기관"].astype(str)))
    yearly = y[["기준기간", "값", "집계기준(판매/등록)", "출처URL", "비고"]].rename(
        columns={"기준기간": "연도", "값": f"{ybasis}(대)", "집계기준(판매/등록)": "집계기준"}).reset_index(drop=True)
    v = yearly[f"{ybasis}(대)"]
    trend = (f"국내 전기차 연도별 {ybasis} — {yearly['연도'].iloc[0]} {int(v.iloc[0]):,}대 → "
             f"{yearly['연도'].iloc[-1]} {int(v.iloc[-1]):,}대")
    # 같은 해 총계가 출처마다 다르면 노트로 남긴다
    other = yearly[yearly["연도"].str.startswith(label) & (v != total)]
    diff_note = (f"같은 해 총계가 다른 출처에서는 {int(other.iloc[0][f'{ybasis}(대)']):,}대다 — 출처마다 집계 범위가 다르다."
                 if len(other) else "")
    return {"label": label, "basis": basis, "org": org, "total": total, "total_url": total_row["출처URL"].iloc[0],
            "makers": makers, "qty": f"{basis}(대)", "yearly": yearly, "yqty": f"{ybasis}(대)",
            "ybasis": ybasis, "yorg": yorg, "trend": trend, "diff_note": diff_note}


def file_hash(path: Path) -> str:
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()[:16]


def _rel(p: Path) -> str:
    p = Path(p)
    p = p if p.is_absolute() else ROOT / p
    return p.resolve().relative_to(ROOT).as_posix()


def load_stamps() -> dict:
    return json.loads(STAMPS.read_text(encoding="utf-8")) if STAMPS.exists() else {}


def stamp(artifact: str, inputs: list) -> None:
    """Claude가 판단해 만든 파일(artifact)이 어떤 입력을 보고 만들어졌는지 기록한다."""
    s = load_stamps()
    s[artifact] = {"입력": {_rel(p): file_hash(ROOT / _rel(p)) for p in inputs},
                   "기록시각": datetime.now().isoformat(timespec="seconds")}
    STAMPS.write_text(json.dumps(s, ensure_ascii=False, indent=2), encoding="utf-8")


def stale_reason(artifact: str, expect: list | None = None) -> str | None:
    """판단 결과가 지금 입력과 맞지 않으면 이유를, 맞으면 None을 돌려준다.

    expect: 이 판단이 반드시 보고 만들었어야 하는 입력. 대상이 바뀌었는지 확인한다.
    """
    if not (ROOT / artifact).exists():
        return "파일 없음"
    rec = load_stamps().get(artifact)
    if not rec:
        return "판단 기록 없음 — 어떤 입력으로 만들었는지 모름"
    for p in expect or []:
        if _rel(p) not in rec["입력"]:
            return f"판단 대상이 다름: {_rel(p)}을(를) 보고 만든 기록이 없음"
    for rel, h in rec["입력"].items():
        p = ROOT / rel
        if not p.exists():
            return f"입력 사라짐: {rel}"
        if file_hash(p) != h:
            return f"입력 바뀜: {rel}"
    return None


VOC_INPUTS = {"아이오닉5": "input/Day2_아이오닉5_VOC리뷰30건_가상.csv", "모델Y": "input/Day2_모델Y_VOC리뷰20건_가상.csv"}


def voc_mismatch(cls: pd.DataFrame) -> str | None:
    """분류표의 리뷰 원문이 지금 input과 같은지 본다. 번호만 같고 내용이 다르면 다른 데이터다(수정 2)."""
    for car, f in VOC_INPUTS.items():
        orig = read_csv(f)[["번호", "리뷰"]]
        m = cls[cls["차종"] == car][["번호", "리뷰"]].merge(orig, on="번호", how="outer", suffixes=("", "_원본"))
        bad = m[m["리뷰"] != m["리뷰_원본"]]
        if len(bad):
            return f"{car} 리뷰 {len(bad)}건이 input과 다름(번호 {bad['번호'].tolist()[:10]})"
    return None


def clinic_config() -> dict:
    """차트 클리닉 대상(원본 CSV · 브랜드 · 원본 이미지). s3가 쓰고 s9가 읽는다(수정 4)."""
    default = {"원본CSV": "input/Day2_차트클리닉_원본데이터_가상.csv", "브랜드": "A사",
               "원본이미지": "input/Day2_차트클리닉_나쁜차트예시_가상.png"}
    if CLINIC_CFG.exists():
        default.update(json.loads(CLINIC_CFG.read_text(encoding="utf-8")))
    return default


# ───────────────────────── xlsx ─────────────────────────
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(color="FFFFFF", bold=True)


def write_xlsx(path: Path, sheets: list[dict]) -> Path:
    """sheets: [{name, df, notes?, formulas?, chart?}]

    formulas: {열이름: "=D{r}/SUM(D$2:D${last})"}  — {r}=행번호, {last}=마지막 데이터 행
    chart: {kind: bar|stacked|line, cat: 열이름, vals: [열이름...], title, y_title?,
            anchor?, colors?: [hex...]}
    """
    wb = Workbook()
    wb.remove(wb.active)
    for s in sheets:
        ws = wb.create_sheet(s["name"][:31])
        df: pd.DataFrame = s["df"]
        cols = list(df.columns) + [c for c in s.get("formulas", {}) if c not in df.columns]
        ws.append(cols)
        for cell in ws[1]:
            cell.fill, cell.font = HEAD_FILL, HEAD_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        n = len(df)
        last = n + 1
        for i, row in enumerate(df.itertuples(index=False), start=2):
            vals = [None if (isinstance(v, float) and pd.isna(v)) else v for v in row]
            for c, tmpl in s.get("formulas", {}).items():
                f = tmpl.format(r=i, last=last)
                if c in df.columns:
                    vals[list(df.columns).index(c)] = f
                else:
                    vals.append(f)
            ws.append(vals)
        for c, fmt in s.get("number_formats", {}).items():
            j = cols.index(c) + 1
            for r in range(2, last + 1):
                ws.cell(r, j).number_format = fmt
        # 열 너비
        for j, c in enumerate(cols, start=1):
            longest = max([len(str(c))] + [len(str(v)) for v in df[c]] if c in df.columns else [len(str(c)), 8])
            ws.column_dimensions[get_column_letter(j)].width = min(max(10, longest * 1.6), 60)
        ws.freeze_panes = "A2"
        # 노트(기준·기간·출처)
        r = last + 2
        for note in filter(None, s.get("notes", [])):
            ws.cell(r, 1, note).font = Font(italic=True, color="555555")
            r += 1
        if s.get("chart"):
            _add_chart(ws, s["chart"], cols, last, r + 1)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    _patch_label_format(path)
    return path


def _patch_label_format(path: Path) -> None:
    """openpyxl은 레이블 numFmt에 sourceLinked를 쓰지 않아 엑셀이 무시한다 — 저장 후 보정."""
    import io
    import zipfile
    with zipfile.ZipFile(path) as zin:
        items = [(i, zin.read(i.filename)) for i in zin.infolist()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item, data in items:
            if item.filename.startswith("xl/charts/chart"):
                data = data.replace(b'<numFmt formatCode="#,##0;-#,##0;;"/>',
                                    b'<numFmt formatCode="#,##0;-#,##0;;" sourceLinked="0"/>')
            zout.writestr(item, data)
    Path(path).write_bytes(buf.getvalue())


def _add_chart(ws, spec: dict, cols: list[str], last: int, below_row: int) -> None:
    kind = spec.get("kind", "bar")
    ch = LineChart() if kind == "line" else BarChart()
    if kind != "line":
        ch.type = "col"
        if kind == "stacked":
            ch.grouping, ch.overlap = "stacked", 100
    ch.title = spec["title"]
    ch.y_axis.number_format = spec.get("y_format", "#,##0")
    ch.y_axis.scaling.min = 0          # 원칙 ②: y축 0부터
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.height, ch.width = 9, 18
    ch.legend.position = "r"
    ci = cols.index(spec["cat"]) + 1
    for v in spec["vals"]:
        vi = cols.index(v) + 1
        ch.add_data(Reference(ws, min_col=vi, min_row=1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=ci, min_row=2, max_row=last))
    colors = spec.get("colors") or [ACCENT, MUTED, NEU, NEG]
    for k, ser in enumerate(ch.series):
        col = colors[k % len(colors)]
        if kind == "line":
            ser.graphicalProperties.line.solidFill = col
            ser.graphicalProperties.line.width = 28000
            ser.smooth = False              # 곡선 보간은 없는 값을 지어낸다
        else:
            ser.graphicalProperties.solidFill = col
            ser.graphicalProperties.line.solidFill = col
    if kind != "line":
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        for attr in ("showSerName", "showCatName", "showLegendKey", "showPercent"):
            setattr(ch.dataLabels, attr, False)
        ch.dataLabels.numFmt = "#,##0;-#,##0;;"   # 0은 레이블을 숨긴다(_patch_label_format이 보정)
    if len(spec["vals"]) == 1:
        ch.legend = None
    ws.add_chart(ch, spec.get("anchor") or f"A{below_row}")


# ───────────────────────── HTML / SVG ─────────────────────────
def esc(s) -> str:
    return html.escape(str(s))


def fmt(v: float, unit: str = "") -> str:
    if isinstance(v, (int,)) or float(v).is_integer():
        s = f"{int(round(v)):,}"
    else:
        s = f"{v:,.1f}"
    return s + unit


def _nice_max(v: float) -> float:
    import math
    if v <= 0:
        return 1
    p = 10 ** math.floor(math.log10(v))
    for m in (1, 1.2, 1.5, 2, 2.5, 3, 4, 5, 6, 8, 10):
        if m * p >= v:
            return m * p
    return 10 * p


def svg_bar(cats, vals, *, unit="", highlight=None, width=640, height=320, horizontal=False) -> str:
    """순위·크기 비교 막대. y축 0 시작. highlight: 강조할 범주 이름(나머지 회색)."""
    if horizontal:
        return _svg_hbar(cats, vals, unit=unit, highlight=highlight, width=width)
    pl, pr, pt, pb = 56, 16, 24, 48
    W, H = width - pl - pr, height - pt - pb
    vmax = _nice_max(max(vals))
    n = len(cats)
    bw = W / n * 0.62
    out = [f'<svg viewBox="0 0 {width} {height}" role="img" class="chart">']
    for k in range(5):
        y = pt + H - H * k / 4
        out.append(f'<line x1="{pl}" x2="{pl+W}" y1="{y:.1f}" y2="{y:.1f}" class="grid"/>')
        out.append(f'<text x="{pl-8}" y="{y+4:.1f}" class="tick" text-anchor="end">{fmt(vmax*k/4)}</text>')
    for i, (c, v) in enumerate(zip(cats, vals)):
        x = pl + W / n * i + (W / n - bw) / 2
        h = H * v / vmax
        cls = "bar hi" if (highlight is None and i == 0) or c == highlight else "bar"
        out.append(f'<rect x="{x:.1f}" y="{pt+H-h:.1f}" width="{bw:.1f}" height="{h:.1f}" class="{cls}"><title>{esc(c)}: {fmt(v, unit)}</title></rect>')
        out.append(f'<text x="{x+bw/2:.1f}" y="{pt+H-h-6:.1f}" class="lbl" text-anchor="middle">{fmt(v, unit)}</text>')
        out.append(f'<text x="{x+bw/2:.1f}" y="{pt+H+18:.1f}" class="tick" text-anchor="middle">{esc(c)}</text>')
    out.append(f'<line x1="{pl}" x2="{pl+W}" y1="{pt+H}" y2="{pt+H}" class="axis"/>')
    out.append("</svg>")
    return "".join(out)


def _svg_hbar(cats, vals, *, unit="", highlight=None, width=640) -> str:
    row = 30
    pl, pr, pt = 150, 70, 8
    height = pt * 2 + row * len(cats)
    W = width - pl - pr
    vmax = max(vals)
    out = [f'<svg viewBox="0 0 {width} {height}" role="img" class="chart">']
    for i, (c, v) in enumerate(zip(cats, vals)):
        y = pt + row * i
        w = W * v / vmax if vmax else 0
        cls = "bar hi" if c == highlight else "bar"
        out.append(f'<text x="{pl-8}" y="{y+row/2+4:.1f}" class="tick" text-anchor="end">{esc(c)}</text>')
        out.append(f'<rect x="{pl}" y="{y+5}" width="{w:.1f}" height="{row-10}" class="{cls}"><title>{esc(c)}: {fmt(v, unit)}</title></rect>')
        out.append(f'<text x="{pl+w+6:.1f}" y="{y+row/2+4:.1f}" class="lbl">{fmt(v, unit)}</text>')
    out.append(f'<line x1="{pl}" x2="{pl}" y1="{pt}" y2="{height-pt}" class="axis"/>')
    out.append("</svg>")
    return "".join(out)


def svg_stacked(cats, series: dict[str, list[float]], *, colors: dict[str, str] | None = None,
                width=640, height=320, percent=False) -> str:
    """구성비 누적 막대. series: {범례: [범주별 값]}."""
    pl, pr, pt, pb = 56, 16, 34, 48
    W, H = width - pl - pr, height - pt - pb
    totals = [sum(s[i] for s in series.values()) for i in range(len(cats))]
    vmax = 100 if percent else _nice_max(max(totals))
    n = len(cats)
    bw = W / n * 0.56
    colors = colors or {}
    out = [f'<svg viewBox="0 0 {width} {height}" role="img" class="chart">']
    for k in range(5):
        y = pt + H - H * k / 4
        out.append(f'<line x1="{pl}" x2="{pl+W}" y1="{y:.1f}" y2="{y:.1f}" class="grid"/>')
        out.append(f'<text x="{pl-8}" y="{y+4:.1f}" class="tick" text-anchor="end">{fmt(vmax*k/4, "%" if percent else "")}</text>')
    for i, c in enumerate(cats):
        x = pl + W / n * i + (W / n - bw) / 2
        y0 = pt + H
        for name, vals in series.items():
            v = vals[i] / totals[i] * 100 if percent and totals[i] else vals[i]
            h = H * v / vmax
            if h > 0:
                out.append(f'<rect x="{x:.1f}" y="{y0-h:.1f}" width="{bw:.1f}" height="{h:.1f}" fill="#{colors.get(name, MUTED)}" stroke="var(--bg)" stroke-width="1"><title>{esc(c)} · {esc(name)}: {fmt(vals[i])}</title></rect>')
                if h > 14:
                    out.append(f'<text x="{x+bw/2:.1f}" y="{y0-h/2+4:.1f}" class="lbl inv" text-anchor="middle">{fmt(vals[i])}</text>')
            y0 -= h
        out.append(f'<text x="{x+bw/2:.1f}" y="{pt+H+18:.1f}" class="tick" text-anchor="middle">{esc(c)}</text>')
    lx = pl
    for name in series:
        out.append(f'<rect x="{lx}" y="8" width="12" height="12" fill="#{colors.get(name, MUTED)}"/>'
                   f'<text x="{lx+16}" y="18" class="tick">{esc(name)}</text>')
        lx += 22 + 12 * len(name) + 10
    out.append(f'<line x1="{pl}" x2="{pl+W}" y1="{pt+H}" y2="{pt+H}" class="axis"/>')
    out.append("</svg>")
    return "".join(out)


def svg_line(cats, series: dict[str, list[float]], *, unit="", ymin=0.0, ymax=None,
             colors: dict[str, str] | None = None, marks: dict[int, str] | None = None,
             width=640, height=320) -> str:
    """시간 흐름 선 차트. marks: {x 인덱스: 설명} — 기준 변경 같은 단절 지점을 세로선으로 표시."""
    pl, pr, pt, pb = 56, 16, 34, 48
    W, H = width - pl - pr, height - pt - pb
    allv = [v for s in series.values() for v in s if v is not None]
    ymax = ymax if ymax is not None else _nice_max(max(allv))
    n = len(cats)
    X = lambda i: pl + W * (i + 0.5) / n
    Y = lambda v: pt + H - H * (v - ymin) / (ymax - ymin)
    colors = colors or {}
    out = [f'<svg viewBox="0 0 {width} {height}" role="img" class="chart">']
    for k in range(5):
        v = ymin + (ymax - ymin) * k / 4
        out.append(f'<line x1="{pl}" x2="{pl+W}" y1="{Y(v):.1f}" y2="{Y(v):.1f}" class="grid"/>')
        out.append(f'<text x="{pl-8}" y="{Y(v)+4:.1f}" class="tick" text-anchor="end">{fmt(v, unit)}</text>')
    for i, c in enumerate(cats):
        out.append(f'<text x="{X(i):.1f}" y="{pt+H+18:.1f}" class="tick" text-anchor="middle">{esc(c)}</text>')
    for i, label in (marks or {}).items():
        x = (X(i - 1) + X(i)) / 2
        out.append(f'<line x1="{x:.1f}" x2="{x:.1f}" y1="{pt}" y2="{pt+H}" class="mark"/>'
                   f'<text x="{x+4:.1f}" y="{pt+12}" class="tick warn">{esc(label)}</text>')
    lx = pl
    for name, vals in series.items():
        col = colors.get(name, ACCENT)
        segs, cur = [], []
        for i, v in enumerate(vals):
            if v is None:
                if cur:
                    segs.append(cur)
                cur = []
            else:
                cur.append((X(i), Y(v), v, cats[i]))
        if cur:
            segs.append(cur)
        for seg in segs:
            pts = " ".join(f"{x:.1f},{y:.1f}" for x, y, _, _ in seg)
            out.append(f'<polyline points="{pts}" fill="none" stroke="#{col}" stroke-width="2.5"/>')
            for x, y, v, c in seg:
                out.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="3.5" fill="#{col}"><title>{esc(c)} · {esc(name)}: {fmt(v, unit)}</title></circle>')
        out.append(f'<rect x="{lx}" y="10" width="14" height="3" fill="#{col}"/><text x="{lx+18}" y="16" class="tick">{esc(name)}</text>')
        lx += 26 + 12 * len(name) + 10
    out.append(f'<line x1="{pl}" x2="{pl+W}" y1="{pt+H}" y2="{pt+H}" class="axis"/>')
    out.append("</svg>")
    return "".join(out)


def img_data_uri(path: Path) -> str:
    b = Path(path).read_bytes()
    return "data:image/png;base64," + base64.b64encode(b).decode()


def table_html(df: pd.DataFrame, max_rows: int | None = None) -> str:
    d = df if max_rows is None else df.head(max_rows)
    head = "".join(f"<th>{esc(c)}</th>" for c in d.columns)
    rows = "".join("<tr>" + "".join(f"<td>{esc('' if pd.isna(v) else v)}</td>" for v in r) + "</tr>"
                   for r in d.itertuples(index=False))
    return f'<div class="tbl"><table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table></div>'


CSS = """
:root{--bg:#ffffff;--fg:#1d2733;--muted:#5b6673;--line:#d9dee5;--card:#f6f8fb;--accent:#1F5FAD;--bar:#A7B0BC;--warn:#B85C00}
@media (prefers-color-scheme:dark){:root:not([data-theme="light"]){--bg:#12171d;--fg:#e6ebf1;--muted:#9aa5b1;--line:#2a333d;--card:#1a2129;--accent:#6aa5ea;--bar:#5d6773;--warn:#f0a050}}
:root[data-theme="dark"]{--bg:#12171d;--fg:#e6ebf1;--muted:#9aa5b1;--line:#2a333d;--card:#1a2129;--accent:#6aa5ea;--bar:#5d6773;--warn:#f0a050}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--fg);font-family:"Pretendard","Malgun Gothic","Apple SD Gothic Neo",system-ui,sans-serif;line-height:1.6}
main{max-width:1040px;margin:0 auto;padding:24px 16px 64px}
header h1{font-size:1.6rem;margin:0 0 4px}
header p{margin:0;color:var(--muted)}
nav{display:flex;gap:8px;flex-wrap:wrap;margin:20px 0}
nav button{border:1px solid var(--line);background:var(--card);color:var(--fg);padding:6px 14px;border-radius:999px;cursor:pointer;font:inherit}
nav button[aria-selected=true]{background:var(--accent);border-color:var(--accent);color:#fff}
section{display:none}section.on{display:block}
.card{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:16px 18px;margin:14px 0}
.card h2{font-size:1.15rem;margin:0 0 6px}
.card h3{font-size:1rem;margin:12px 0 4px}
.note{font-size:.85rem;color:var(--muted)}
.tag{display:inline-block;font-size:.75rem;padding:1px 8px;border-radius:999px;border:1px solid var(--warn);color:var(--warn);margin-left:6px;vertical-align:middle}
.grid2{display:grid;grid-template-columns:1fr 1fr;gap:14px}
@media (max-width:760px){.grid2{grid-template-columns:1fr}}
svg.chart{width:100%;height:auto;display:block}
svg .grid{stroke:var(--line);stroke-width:1}
svg .axis{stroke:var(--muted);stroke-width:1}
svg .mark{stroke:var(--warn);stroke-width:1.5;stroke-dasharray:4 3}
svg .tick{fill:var(--muted);font-size:12px}
svg .warn{fill:var(--warn)}
svg .lbl{fill:var(--fg);font-size:12px;font-weight:600}
svg .lbl.inv{fill:#fff}
svg .bar{fill:var(--bar)} svg .bar.hi{fill:var(--accent)}
img.full{width:100%;height:auto;border-radius:8px;border:1px solid var(--line)}
.tbl{overflow-x:auto}
table{border-collapse:collapse;width:100%;font-size:.88rem}
th,td{border-bottom:1px solid var(--line);padding:6px 8px;text-align:left;vertical-align:top}
th{color:var(--muted);font-weight:600}
"""

TAB_JS = """
document.querySelectorAll('nav button').forEach(b=>b.addEventListener('click',()=>{
 document.querySelectorAll('nav button').forEach(x=>x.setAttribute('aria-selected',x===b));
 document.querySelectorAll('section').forEach(s=>s.classList.toggle('on',s.id===b.dataset.t));
 try{localStorage.setItem('d2tab',b.dataset.t)}catch(e){}
}));
try{const t=localStorage.getItem('d2tab');const b=t&&document.querySelector('nav button[data-t="'+t+'"]');if(b)b.click()}catch(e){}
"""


def write_html(path: Path, title: str, subtitle: str, tabs: list[tuple[str, str, str]]) -> Path:
    """tabs: [(id, 탭이름, 본문 HTML)] — 단일 파일, 외부 리소스 없음."""
    nav = "".join(f'<button data-t="{t}" aria-selected="{str(i==0).lower()}">{esc(n)}</button>'
                  for i, (t, n, _) in enumerate(tabs))
    secs = "".join(f'<section id="{t}" class="{"on" if i==0 else ""}">{body}</section>'
                   for i, (t, _, body) in enumerate(tabs))
    doc = (f'<!doctype html><html lang="ko"><head><meta charset="utf-8">'
           f'<meta name="viewport" content="width=device-width,initial-scale=1">'
           f'<title>{esc(title)}</title><style>{CSS}</style></head><body><main>'
           f'<header><h1>{esc(title)}</h1><p>{esc(subtitle)}</p></header>'
           f'<nav role="tablist">{nav}</nav>{secs}</main><script>{TAB_JS}</script></body></html>')
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(doc, encoding="utf-8")
    return path
